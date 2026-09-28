from django.contrib import admin
from django.contrib.auth.admin import UserAdmin as BaseUserAdmin
from django.contrib.auth.models import User
from django.db.models import Count
from django.utils.html import format_html
from django.utils import timezone
from .models import Role, UserProfile, Store, EnquiryType, Grievance


# ---------------------------------------------------------
# User & Profile Admin
# ---------------------------------------------------------

class UserProfileInline(admin.StackedInline):
    model = UserProfile
    can_delete = False
    verbose_name_plural = 'Profile'
    filter_horizontal = ('stores',)


class CustomUserAdmin(BaseUserAdmin):
    inlines = (UserProfileInline,)
    list_display = ('username', 'email', 'get_emp_id', 'get_role', 'get_status', 'is_staff', 'is_superuser')
    
    def get_emp_id(self, obj):
        return obj.profile.emp_id if hasattr(obj, 'profile') else '-'
    get_emp_id.short_description = 'Emp ID'
    
    def get_role(self, obj):
        return obj.profile.role.name if hasattr(obj, 'profile') and obj.profile.role else '-'
    get_role.short_description = 'Role'
    
    def get_status(self, obj):
        return obj.profile.get_approval_status_display() if hasattr(obj, 'profile') else '-'
    get_status.short_description = 'Approval Status'


admin.site.unregister(User)
admin.site.register(User, CustomUserAdmin)


@admin.register(UserProfile)
class UserProfileAdmin(admin.ModelAdmin):
    list_display = ('user', 'emp_id', 'role', 'approval_status', 'is_approved', 'created_at')
    list_filter = ('role', 'approval_status', 'is_approved')
    search_fields = ('user__username', 'emp_id', 'user__email')
    filter_horizontal = ('stores',)


# ---------------------------------------------------------
# Master Data Admin: Role, Store, EnquiryType
# ---------------------------------------------------------

@admin.register(Role)
class RoleAdmin(admin.ModelAdmin):
    list_display = ('name', 'description', 'user_count', 'created_at')
    search_fields = ('name', 'description')
    ordering = ('name',)

    def get_queryset(self, request):
        qs = super().get_queryset(request)
        return qs.annotate(users_total=Count('users'))

    @admin.display(description='Assigned Users')
    def user_count(self, obj):
        return obj.users_total


@admin.register(Store)
class StoreAdmin(admin.ModelAdmin):
    list_display = ('name', 'is_active', 'grievance_count', 'created_at')
    list_editable = ('is_active',)
    search_fields = ('name',)
    list_filter = ('is_active',)
    actions = ['deactivate_selected', 'activate_selected']

    def get_queryset(self, request):
        qs = super().get_queryset(request)
        return qs.annotate(grievance_total=Count('grievances'))

    @admin.display(description='Total Grievances')
    def grievance_count(self, obj):
        return obj.grievance_total

    @admin.action(description='Deactivate selected stores')
    def deactivate_selected(self, request, queryset):
        count = queryset.update(is_active=False)
        self.message_user(request, f"{count} store(s) successfully deactivated.")

    @admin.action(description='Activate selected stores')
    def activate_selected(self, request, queryset):
        count = queryset.update(is_active=True)
        self.message_user(request, f"{count} store(s) successfully activated.")


@admin.register(EnquiryType)
class EnquiryTypeAdmin(admin.ModelAdmin):
    list_display = ('name', 'is_active', 'grievance_count', 'created_at')
    list_editable = ('is_active',)
    search_fields = ('name',)
    list_filter = ('is_active',)
    actions = ['deactivate_selected', 'activate_selected']

    def get_queryset(self, request):
        qs = super().get_queryset(request)
        return qs.annotate(grievance_total=Count('grievances'))

    @admin.display(description='Total Grievances')
    def grievance_count(self, obj):
        return obj.grievance_total

    @admin.action(description='Deactivate selected enquiry types')
    def deactivate_selected(self, request, queryset):
        count = queryset.update(is_active=False)
        self.message_user(request, f"{count} enquiry type(s) successfully deactivated.")

    @admin.action(description='Activate selected enquiry types')
    def activate_selected(self, request, queryset):
        count = queryset.update(is_active=True)
        self.message_user(request, f"{count} enquiry type(s) successfully activated.")


# ---------------------------------------------------------
# Grievance Admin with Excel Export & Status Badges
# ---------------------------------------------------------

@admin.register(Grievance)
class GrievanceAdmin(admin.ModelAdmin):
    list_display = (
        'id', 'open_date', 'emp_id', 'emp_name', 'store', 'enquiry_type', 
        'enquiry_received_by', 'status_badge', 'closed_date', 'period_days'
    )
    list_filter = ('current_status', 'store', 'enquiry_type', 'open_date', 'closed_date')
    search_fields = ('id', 'emp_id', 'emp_name', 'enquiry_details', 'closing_reason', 'enquiry_received_by')
    list_select_related = ('store', 'enquiry_type', 'created_by')
    date_hierarchy = 'open_date'
    actions = ['export_to_excel', 'mark_as_closed']
    readonly_fields = ('period_days', 'created_at', 'updated_at')

    fieldsets = (
        ('Employee & Store Information', {
            'fields': (
                ('open_date', 'current_status'),
                ('emp_id', 'emp_name'),
                ('designation', 'department'),
                ('section', 'store'),
            )
        }),
        ('Enquiry Details', {
            'fields': (
                ('enquiry_type', 'enquiry_received_by'),
                'enquiry_details',
            )
        }),
        ('Resolution Workflow & Closure', {
            'fields': (
                ('first_level', 'second_level'),
                ('closed_date', 'period_days'),
                'closing_reason',
            )
        }),
        ('Audit Information', {
            'classes': ('collapse',),
            'fields': (
                'created_by',
                ('created_at', 'updated_at'),
            )
        }),
    )

    def save_model(self, request, obj, form, change):
        if not change and not obj.created_by:
            obj.created_by = request.user
        super().save_model(request, obj, form, change)

    @admin.display(description='Status')
    def status_badge(self, obj):
        colors = {
            'open': '#dc3545',
            'hold': '#fd7e14',
            'extension': '#0d6efd',
            'closed': '#198754',
        }
        color = colors.get(obj.current_status, '#6c757d')
        return format_html(
            '<span style="background-color: {}; color: white; padding: 3px 8px; border-radius: 4px; font-weight: 600; font-size: 11px; text-transform: uppercase;">{}</span>',
            color,
            obj.get_current_status_display()
        )

    @admin.action(description='Export Selected Grievances to Excel (.xlsx)')
    def export_to_excel(self, request, queryset):
        import openpyxl
        from django.http import HttpResponse

        queryset = queryset.select_related('store', 'enquiry_type')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "GRIEVANCES - RRPL"

        headers = [
            "ID", "Open Date", "Emp ID", "Emp Name", "Designation",
            "Department", "Section", "Store Name", "Enquiry Type",
            "Enquiry Received By", "Enquiry Details", "1st Level",
            "2nd Level", "Closed Date", "Period Days", "Closing Reason", "Current Status"
        ]
        ws.append(headers)

        header_font = openpyxl.styles.Font(bold=True, color="FFFFFF", size=11)
        header_fill = openpyxl.styles.PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_align = openpyxl.styles.Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        ws.row_dimensions[1].height = 25

        for obj in queryset:
            ws.append([
                obj.id,
                obj.open_date.strftime('%Y-%m-%d') if obj.open_date else '',
                obj.emp_id,
                obj.emp_name,
                obj.designation,
                obj.department,
                obj.section,
                obj.store.name if obj.store else '',
                obj.enquiry_type.name if obj.enquiry_type else '',
                obj.enquiry_received_by,
                obj.enquiry_details,
                obj.first_level or '',
                obj.second_level or '',
                obj.closed_date.strftime('%Y-%m-%d') if obj.closed_date else '',
                obj.period_days if obj.period_days is not None else '',
                obj.closing_reason or '',
                obj.get_current_status_display(),
            ])

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"Grievance_Report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response

    @admin.action(description='Mark Selected as Closed')
    def mark_as_closed(self, request, queryset):
        today = timezone.now().date()
        count = 0
        for g in queryset:
            g.current_status = 'closed'
            if not g.closed_date:
                g.closed_date = today
            if g.open_date:
                g.period_days = max(0, (g.closed_date - g.open_date).days)
            g.save()
            count += 1
        self.message_user(request, f"{count} grievance(s) successfully marked as Closed.")
